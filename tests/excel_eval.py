"""A small evaluator for the Excel formulas the workbook writes, so the tests can
calculate a workbook without a spreadsheet program.

It covers only what ``marketing_roi.workbook`` uses: numbers, strings, TRUE/FALSE, cell
and range references (optionally on another sheet, with or without $), the operators
+ - * / & = <> < <= > >= and unary minus, and SUM, COUNT, MIN, MAX, SUMIFS, COUNTIFS,
IF, OR, NOT, DATE, YEAR and MONTH, with Excel's rules where they matter here: text
compares without case, an empty cell is 0 in arithmetic, DATE rolls months over, and a
criterion such as "<=45000" compares numbers. Anything else raises, so a formula the
evaluator does not know cannot pass as a wrong number.
"""

from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

EPOCH = dt.date(1899, 12, 30)

TOKEN = re.compile(
    r"""\s*(?:
      (?P<str>"(?:[^"]|"")*")
    | (?P<num>\d+(?:\.\d*)?(?:[eE][-+]?\d+)?)
    | (?P<ref>(?:(?P<sheet>[A-Za-z_][\w.]*)!)?\$?(?P<c1>[A-Z]{1,3})\$?(?P<r1>\d+)
              (?::\$?(?P<c2>[A-Z]{1,3})\$?(?P<r2>\d+))?)
    | (?P<func>[A-Z]+)\(
    | (?P<bool>TRUE|FALSE)\b
    | (?P<op><>|<=|>=|[-+*/&=<>(),])
    )""",
    re.VERBOSE,
)


def tokens(formula: str) -> list[tuple[str, object]]:
    out, i = [], 0
    text = formula.strip()
    while i < len(text):
        m = TOKEN.match(text, i)
        if not m or m.end() == i:
            raise ValueError(f"cannot read {text[i:]!r} in {formula!r}")
        i = m.end()
        kind = m.lastgroup if m.lastgroup not in ("sheet", "c1", "r1", "c2", "r2") else "ref"
        if m.group("ref"):
            kind = "ref"
        if kind == "str":
            out.append(("str", m.group("str")[1:-1].replace('""', '"')))
        elif kind == "num":
            out.append(("num", float(m.group("num"))))
        elif kind == "ref":
            c1, r1 = column_index_from_string(m.group("c1")), int(m.group("r1"))
            c2 = column_index_from_string(m.group("c2")) if m.group("c2") else c1
            r2 = int(m.group("r2")) if m.group("r2") else r1
            out.append(("ref", (m.group("sheet"), c1, r1, c2, r2, bool(m.group("c2")))))
        elif kind == "func":
            out.append(("func", m.group("func")))
        elif kind == "bool":
            out.append(("num", m.group("bool") == "TRUE"))
        else:
            out.append(("op", m.group("op")))
    return out


def _num(v) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, bool):
        return float(v)
    if isinstance(v, int | float):
        return float(v)
    raise TypeError(f"#VALUE!: {v!r} is not a number")


def _compare(op: str, a, b) -> bool:
    a = "" if a is None else a
    b = "" if b is None else b
    if isinstance(a, str) and isinstance(b, str):
        a, b = a.lower(), b.lower()
    elif isinstance(a, str) or isinstance(b, str):
        # Excel ranks any text above any number; only = and <> are used across types here.
        if op not in ("=", "<>"):
            raise TypeError(f"comparing {a!r} {op} {b!r}")
        return op == "<>"
    else:
        a, b = _num(a), _num(b)
    return {
        "=": a == b, "<>": a != b, "<": a < b, "<=": a <= b, ">": a > b, ">=": a >= b,
    }[op]  # fmt: skip


def _matches(cell, criterion) -> bool:
    if isinstance(criterion, str):
        m = re.match(r"(<>|<=|>=|<|>|=)(.*)$", criterion)
        if m:
            op, operand = m.groups()
            try:
                value = float(operand)
            except ValueError:
                return _compare(op, cell, operand)
            if not isinstance(cell, int | float) or isinstance(cell, bool) or cell is None:
                return op == "<>"
            return _compare(op, cell, value)
        return isinstance(cell, str) and cell.lower() == criterion.lower()
    if isinstance(cell, int | float) and not isinstance(cell, bool):
        return float(cell) == _num(criterion)
    return False


def _date(y, m, d) -> float:
    y, m = int(y) + (int(m) - 1) // 12, (int(m) - 1) % 12 + 1
    return float((dt.date(y, m, 1) - EPOCH).days + int(d) - 1)


def _from_serial(v) -> dt.date:
    return EPOCH + dt.timedelta(days=int(_num(v)))


class Book:
    """A workbook's cells, calculated on demand."""

    def __init__(self, path: Path):
        wb = load_workbook(path)
        self.order = wb.sheetnames
        self.cells: dict[str, dict[tuple[int, int], object]] = {}
        for ws in wb.worksheets:
            cells = {}
            for row in ws.iter_rows():
                for c in row:
                    v = c.value
                    if isinstance(v, dt.datetime):
                        v = float((v.date() - EPOCH).days)
                    if v is not None:
                        cells[(c.row, c.column)] = v
            self.cells[ws.title] = cells
        self.done: dict[tuple[str, int, int], object] = {}
        self.ranges: dict[tuple, list] = {}

    def set(self, sheet: str, row: int, col: int, value) -> None:
        """Change an input cell; every calculated value is dropped."""
        self.cells[sheet][(row, col)] = value
        self.done.clear()
        self.ranges.clear()

    def calculate(self, first: tuple[str, ...] = ("Daily", "Purchases", "About")) -> None:
        """Every formula, in an order that keeps the chains down a column short."""
        for sheet in [*first, *(s for s in self.order if s not in first)]:
            for row, col in sorted(self.cells[sheet]):
                self.value(sheet, row, col)

    def value(self, sheet: str, row: int, col: int):
        key = (sheet, row, col)
        if key in self.done:
            return self.done[key]
        raw = self.cells[sheet].get((row, col))
        if isinstance(raw, str) and raw.startswith("="):
            result = _Parser(self, sheet, tokens(raw[1:])).run()
        else:
            result = raw
        self.done[key] = result
        return result

    def range(self, sheet: str, c1: int, r1: int, c2: int, r2: int) -> list:
        key = (sheet, c1, r1, c2, r2)
        if key not in self.ranges:
            if c1 != c2:
                raise ValueError("only single-column ranges are used")
            self.ranges[key] = [self.value(sheet, r, c1) for r in range(r1, r2 + 1)]
        return self.ranges[key]


class _Parser:
    def __init__(self, book: Book, sheet: str, toks: list):
        self.book, self.sheet, self.toks, self.i = book, sheet, toks, 0

    def run(self):
        v = self.compare()
        if self.i != len(self.toks):
            raise ValueError(f"unexpected {self.toks[self.i]!r}")
        return v

    def peek(self):
        return self.toks[self.i] if self.i < len(self.toks) else (None, None)

    def take(self, op: str) -> None:
        if self.peek() != ("op", op):
            raise ValueError(f"expected {op!r}, got {self.peek()!r}")
        self.i += 1

    def compare(self):
        a = self.concat()
        while self.peek()[0] == "op" and self.peek()[1] in ("=", "<>", "<", "<=", ">", ">="):
            op = self.toks[self.i][1]
            self.i += 1
            a = _compare(op, a, self.concat())
        return a

    def concat(self):
        a = self.additive()
        while self.peek() == ("op", "&"):
            self.i += 1
            b = self.additive()
            a = _text(a) + _text(b)
        return a

    def additive(self):
        a = self.term()
        while self.peek()[0] == "op" and self.peek()[1] in "+-":
            op = self.toks[self.i][1]
            self.i += 1
            b = self.term()
            a = _num(a) + _num(b) if op == "+" else _num(a) - _num(b)
        return a

    def term(self):
        a = self.unary()
        while self.peek()[0] == "op" and self.peek()[1] in "*/":
            op = self.toks[self.i][1]
            self.i += 1
            b = self.unary()
            if op == "*":
                a = _num(a) * _num(b)
            else:
                if _num(b) == 0:
                    raise ZeroDivisionError("#DIV/0!")
                a = _num(a) / _num(b)
        return a

    def unary(self):
        if self.peek() == ("op", "-"):
            self.i += 1
            return -_num(self.unary())
        return self.primary()

    def primary(self):
        kind, v = self.peek()
        self.i += 1
        if kind in ("num", "str"):
            return v
        if kind == "ref":
            sheet, c1, r1, c2, r2, is_range = v
            sheet = sheet or self.sheet
            if is_range:
                return ("range", self.book.range(sheet, c1, r1, c2, r2))
            return self.book.value(sheet, r1, c1)
        if kind == "op" and v == "(":
            out = self.compare()
            self.take(")")
            return out
        if kind == "func":
            return self.call(v)
        raise ValueError(f"unexpected {kind} {v!r}")

    def args(self, lazy: bool = False) -> list:
        """The arguments up to the closing bracket; with ``lazy``, unevaluated spans."""
        out = []
        if self.peek() == ("op", ")"):
            self.i += 1
            return out
        while True:
            if lazy:
                start, depth = self.i, 0
                while not (depth == 0 and self.peek()[0] == "op" and self.peek()[1] in ",)"):
                    k, x = self.peek()
                    if k == "func" or (k == "op" and x == "("):
                        depth += 1
                    elif k == "op" and x == ")":
                        depth -= 1
                    self.i += 1
                out.append(self.toks[start : self.i])
            else:
                out.append(self.compare())
            if self.peek() == ("op", ","):
                self.i += 1
                continue
            self.take(")")
            return out

    def call(self, name: str):
        if name == "IF":
            cond, yes, no = self.args(lazy=True)
            chosen = yes if self.sub(cond) else no
            return self.sub(chosen)
        a = self.args()
        values = lambda x: x[1] if isinstance(x, tuple) and x[0] == "range" else [x]  # noqa: E731
        if name == "SUM":
            return sum(_num(v) for x in a for v in values(x) if not isinstance(v, str))
        if name == "COUNT":
            return float(sum(isinstance(v, int | float) and not isinstance(v, bool)
                             for x in a for v in values(x)))  # fmt: skip
        if name in ("MIN", "MAX"):
            nums = [_num(v) for x in a for v in values(x) if isinstance(v, int | float)]
            return (min if name == "MIN" else max)(nums) if nums else 0.0
        if name in ("SUMIFS", "COUNTIFS"):
            total = values(a[0]) if name == "SUMIFS" else None
            pairs = a[1:] if name == "SUMIFS" else a
            if len(pairs) % 2:
                raise ValueError(f"{name} needs range and criterion pairs")
            rows = None
            for rng, crit in zip(pairs[::2], pairs[1::2], strict=True):
                cells = values(rng)
                if rows is None:
                    rows = list(range(len(cells)))
                rows = [i for i in rows if _matches(cells[i], crit)]
            if name == "COUNTIFS":
                return float(len(rows))
            return sum(_num(total[i]) for i in rows if not isinstance(total[i], str))
        if name == "OR":
            return any(_truth(v) for x in a for v in values(x))
        if name == "NOT":
            return not _truth(a[0])
        if name == "DATE":
            return _date(*(_num(x) for x in a))
        if name == "YEAR":
            return float(_from_serial(a[0]).year)
        if name == "MONTH":
            return float(_from_serial(a[0]).month)
        raise ValueError(f"function {name} is not supported")

    def sub(self, toks: list):
        return _Parser(self.book, self.sheet, toks).run()


def _truth(v) -> bool:
    if isinstance(v, str):
        if v == "":
            return False
        raise TypeError(f"#VALUE!: {v!r} is not TRUE or FALSE")
    return bool(_num(v))


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)
