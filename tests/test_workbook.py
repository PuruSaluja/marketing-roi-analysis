"""The Excel workbook, calculated by ``excel_eval`` and checked figure by figure against
the analysis views it rebuilds."""

import datetime as dt
from decimal import Decimal

import psycopg
import pytest
from excel_eval import Book
from openpyxl import load_workbook

from marketing_roi import workbook
from marketing_roi.analysis import create_views
from marketing_roi.generate import GeneratorConfig, generate, write
from marketing_roi.load import load

# Crosses New Year and ends on 2026-01-30, so the last month is incomplete.
CONFIG = GeneratorConfig(seed=5, start=dt.date(2025, 11, 1), days=91)


@pytest.fixture(scope="module")
def loaded(test_dsn, tmp_path_factory) -> str:
    out = tmp_path_factory.mktemp("generated")
    write(generate(CONFIG), out)
    with psycopg.connect(test_dsn, autocommit=True) as conn:
        conn.execute("DROP SCHEMA IF EXISTS analysis, marketing, truth CASCADE")
    load(test_dsn, out)
    create_views(test_dsn)
    return test_dsn


@pytest.fixture(scope="module")
def built(loaded, tmp_path_factory):
    path = tmp_path_factory.mktemp("workbook") / "model.xlsx"
    inputs = workbook.read_inputs(loaded)
    # A 30-day window: the run's 91 days follow no customer for the default 180.
    rows = workbook.build(inputs, path, ltv_days=30)
    return path, inputs, rows


def headers_of(book: Book) -> dict[str, list[str]]:
    out = {}
    for name in [*workbook.SHEETS, "LTV"]:
        top = 4 if name == "LTV" else 1
        cells = book.cells[name]
        out[name] = [cells[(top, c)] for c in range(1, 1 + max(c for r, c in cells if r == top))]
    return out


def calculated(path):
    book = Book(path)
    book.calculate()
    return book, headers_of(book), lambda s, r, c: book.value(s, r, c)


@pytest.fixture(scope="module")
def calc(built):
    """The workbook calculated once, for the tests that only read it."""
    return calculated(built[0])


def test_every_figure_matches_the_view_it_rebuilds(loaded, built, calc) -> None:
    path, inputs, rows = built
    book, headers, value = calc
    result = workbook.check(loaded, headers, value)
    assert result["wrong"] == []
    with psycopg.connect(loaded) as conn:
        ltv = conn.execute("SELECT count(*) FROM analysis.ltv_by_channel(30)").fetchone()[0]
    assert ltv >= 4
    # Figures per row: Channels 4 channels x (3 months + all) x 11, Acquisition 7 groups x
    # 4 x 4, Blended 4 x 11, Cohorts 3 + 2 + 1 rows x 9; LTV 6 per group in the view, and
    # 1 (customers, 0) per group it leaves out.
    assert result["compared"] == 176 + 112 + 44 + 54 + 6 * ltv + (7 - ltv)
    assert rows["Daily"] == len(inputs.daily) + 1 == 4 * 91 + 1
    assert rows["Purchases"] == len(inputs.purchases) + 1


def test_the_first_purchase_flag_follows_the_view_s_order(loaded, built, calc) -> None:
    path, inputs, _ = built
    book, _, _ = calc
    with psycopg.connect(loaded) as conn:
        first = dict(
            conn.execute("SELECT session_id, first_purchase FROM analysis.purchases").fetchall()
        )
    flags = {book.value("Purchases", r, 2): book.value("Purchases", r, 8)
             for r in range(2, len(inputs.purchases) + 2)}  # fmt: skip
    assert flags == {k: float(v) for k, v in first.items()}
    assert sum(flags.values()) < len(flags)  # some customers bought again


def test_the_ltv_window_is_an_input_the_sheet_follows(loaded, built) -> None:
    path, _, _ = built
    book, headers, value = calculated(path)
    at_30 = [book.value("LTV", r, 6) for r in range(5, 12)]
    book.set("LTV", 1, 2, 60)
    book.calculate()
    assert workbook.check(loaded, headers, value)["wrong"] == []
    assert [book.value("LTV", r, 6) for r in range(5, 12)] != at_30


def test_a_wrong_formula_is_found(loaded, built) -> None:
    path, _, _ = built
    book, headers, value = calculated(path)
    book.set("Channels", 2, 4, "=Daily!$E$2*2")  # a month's spend, made up
    book.calculate()
    wrong = workbook.check(loaded, headers, value)["wrong"]
    # The spend, and the ratios on the same row that divide by it or into it.
    assert {(s, r) for s, r, *_ in wrong} == {("Channels", 2)}
    assert {k for *_, k, _, _ in wrong} >= {"spend", "cpc", "roas"}


def test_a_figure_in_a_row_the_view_leaves_out_is_found(loaded, built) -> None:
    path, _, _ = built
    book, headers, value = calculated(path)
    book.set("LTV", 1, 2, 91)  # only customers from the first day are followed for 91
    book.calculate()
    with psycopg.connect(loaded) as conn:
        kept = {r[0] for r in conn.execute("SELECT channel_group FROM analysis.ltv_by_channel(91)")}
    left_out = [r for r in range(5, 12) if book.value("LTV", r, 1) not in kept]
    assert left_out and all(book.value("LTV", r, 3) == 0 for r in left_out)
    assert workbook.check(loaded, headers, value)["wrong"] == []
    book.set("LTV", left_out[0], 3, 5)
    wrong = workbook.check(loaded, headers, value)["wrong"]
    assert wrong == [("LTV", left_out[0], "customers", 5, "row not in the view")]


def test_a_row_missing_from_a_sheet_is_found(loaded, built) -> None:
    path, _, _ = built
    book, headers, value = calculated(path)
    last = max(r for r, _ in book.cells["Blended"])
    for col in range(1, 13):
        book.set("Blended", last, col, None)  # the whole period's row
    wrong = workbook.check(loaded, headers, value)["wrong"]
    assert wrong == [("Blended", None, (None,), None, "view row missing from the sheet")]


def test_a_calculated_and_saved_workbook_is_checked_from_its_values(
    loaded, built, calc, tmp_path
) -> None:
    path, _, _ = built
    book, _, _ = calc
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    c.value = book.value(ws.title, c.row, c.column)
    saved = tmp_path / "calculated.xlsx"
    wb.save(saved)
    assert workbook.check_file(loaded, saved)["wrong"] == []
    assert workbook.main(["--dsn", loaded, "--check", str(saved)]) == 0


def test_formulas_without_values_are_not_taken_for_figures(loaded, built, capsys) -> None:
    path, _, _ = built
    with pytest.raises(workbook.WorkbookError, match="without calculated values"):
        workbook.check_file(loaded, path)
    assert workbook.main(["--dsn", loaded, "--check", str(path)]) == 1
    assert "without calculated values" in capsys.readouterr().err


def test_the_command_line_writes_the_workbook(loaded, tmp_path, capsys) -> None:
    out = tmp_path / "out" / "model.xlsx"
    assert workbook.main(["--dsn", loaded, "--out", str(out), "--ltv-days", "60"]) == 0
    assert "wrote" in capsys.readouterr().out
    assert load_workbook(out)["LTV"]["B1"].value == 60


@pytest.mark.parametrize("bad", [0, -1, 1.5, True])
def test_a_bad_window_is_refused(built, tmp_path, bad) -> None:
    _, inputs, _ = built
    with pytest.raises(workbook.WorkbookError, match="ltv_days"):
        workbook.build(inputs, tmp_path / "x.xlsx", ltv_days=bad)


def test_the_command_line_needs_a_database(monkeypatch, capsys) -> None:
    monkeypatch.delenv("MARKETING_ROI_DSN", raising=False)
    with pytest.raises(SystemExit):
        workbook.main(["--out", "x.xlsx"])
    assert "no database" in capsys.readouterr().err


def test_a_database_without_the_views_is_refused(test_dsn) -> None:
    # A database that does not exist, so the loaded one the other tests read stays.
    with pytest.raises(workbook.WorkbookError, match="cannot read the analysis views"):
        workbook.read_inputs(f"{test_dsn} dbname=marketing_roi_no_such_test")


def test_serial_days_are_excel_s() -> None:
    assert workbook.serial(dt.date(2025, 1, 1)) == 45658
    assert workbook.from_serial(45658) == dt.date(2025, 1, 1)
    assert workbook.from_serial(dt.datetime(2025, 1, 1)) == dt.date(2025, 1, 1)


def small_inputs(name: str = "paid_x", user: str = "1.2"):
    day = dt.date(2025, 1, 1)
    return workbook.Inputs(
        daily=[dict(date=day, channel=name, channel_group="Paid X", spend=Decimal("10"),
                    impressions=5, clicks=1, sessions=1, purchases=1, revenue=Decimal("20"))],
        purchases=[dict(user_pseudo_id=user, session_id=1, date=day, channel_group="Paid X",
                        paid=True, revenue=Decimal("20"))],
    )  # fmt: skip


def test_input_text_that_looks_like_a_formula_stays_text(tmp_path) -> None:
    path = tmp_path / "text.xlsx"
    workbook.build(small_inputs(name="paid_x", user='=HYPERLINK("http://example.com")'), path)
    wb = load_workbook(path)
    daily, purchases = wb["Daily"], wb["Purchases"]
    assert purchases["A2"].data_type == "s"
    assert purchases["A2"].value == '=HYPERLINK("http://example.com")'
    assert daily["B2"].data_type == "f" and purchases["H2"].data_type == "f"  # still formulas
    assert daily["C2"].data_type == "s"


@pytest.mark.parametrize("name", ["=1+1", "<10", ">x", "Paid*", "Paid?", "a~b"])
def test_names_a_criterion_would_not_match_as_written_are_refused(tmp_path, name) -> None:
    with pytest.raises(workbook.WorkbookError, match="would not match as written"):
        workbook.build(small_inputs(name=name, user="1.2"), tmp_path / "x.xlsx")


def test_more_rows_than_a_sheet_holds_are_refused(monkeypatch, tmp_path) -> None:
    monkeypatch.setattr(workbook, "MAX_ROWS", 1)  # the header alone fills it
    with pytest.raises(workbook.WorkbookError, match="do not fit a sheet's 1 rows"):
        workbook.build(small_inputs(), tmp_path / "big.xlsx")
    assert not (tmp_path / "big.xlsx").exists()


def test_a_workbook_that_cannot_be_written_is_refused(tmp_path) -> None:
    with pytest.raises(workbook.WorkbookError, match="cannot write"):
        workbook.build(small_inputs(), tmp_path)  # a directory


@pytest.mark.parametrize("name, content", [("missing.xlsx", None), ("damaged.xlsx", "not a zip"),
                                           ("notes.md", "# notes")])  # fmt: skip
def test_a_file_that_is_not_a_workbook_is_refused_plainly(tmp_path, capsys, name, content) -> None:
    path = tmp_path / name
    if content is not None:
        path.write_text(content, encoding="utf-8")
    with pytest.raises(workbook.WorkbookError, match="cannot read .* as a workbook"):
        workbook.check_file("dbname=unused", path)
    assert workbook.main(["--dsn", "dbname=unused", "--check", str(path)]) == 1
    assert "as a workbook" in capsys.readouterr().err


@pytest.mark.parametrize("window", [30.5, float("inf"), 0, "30", True])
def test_a_window_that_is_not_a_whole_number_of_days_is_refused(loaded, built, window) -> None:
    # The window is read before any figure, so nothing needs calculating.
    headers = headers_of(Book(built[0]))

    def value(sheet, row, col):
        assert (sheet, row, col) == ("LTV", 1, 2)
        return window

    with pytest.raises(workbook.WorkbookError, match="whole number of days"):
        workbook.check(loaded, headers, value)
