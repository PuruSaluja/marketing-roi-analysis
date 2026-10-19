"""An Excel workbook that rebuilds the analysis views' figures in formulas.

``python -m marketing_roi.workbook --out model.xlsx`` reads two of the ``analysis``
views' rows from PostgreSQL as inputs and writes a workbook in which every figure of
``channel_performance``, ``acquisition``, ``blended``, ``cohorts`` and
``ltv_by_channel`` is an Excel formula over them, so the logic can be read, changed and
followed in a spreadsheet:

``Daily``
    ``analysis.channel_daily`` as values: each paid channel's spend, impressions and
    clicks a day, and the sessions, purchases and revenue GA4 credits to it. The
    sessions are summed per day in SQL because there are more of them (2.4 million on
    the default run) than a sheet's 1,048,576 rows.
``Purchases``
    Every purchase (``analysis.purchases``) as values, sorted by customer and time, with
    the rest in formulas: whether it is the customer's first (the row above is another
    customer, as ``row_number() = 1`` over the same order), the customer's first day,
    cohort month and acquiring channel group carried down from that row, the months and
    days since, and whether it is the customer's first purchase that month.
``Channels``, ``Acquisition``, ``Blended``, ``Cohorts``, ``LTV``
    One sheet per view, row for row, in SUMIFS and COUNTIFS over the two input sheets.
    An empty cell is SQL's NULL (a ratio with nothing to divide by). ``LTV``'s window
    (180 days) is an input cell: change it and the sheet follows.

openpyxl writes formulas without their results. ``--check`` reads a workbook that a
spreadsheet program has calculated and saved, and compares every figure with the views;
on Windows, ``scripts/excel_recalc.ps1`` has Excel do the calculating.
"""

from __future__ import annotations

import argparse
import datetime as dt
import math
import os
import sys
from collections.abc import Callable, Iterable
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

import psycopg
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from psycopg.rows import dict_row

EPOCH = dt.date(1899, 12, 30)
"""Day 0 of Excel's 1900 date system (as Excel counts it, past its 1900 leap-year bug)."""
DATE_FORMAT = "yyyy-mm-dd"
LTV_DAYS = 180

DAILY = ("date", "month", "channel", "channel_group", "spend", "impressions", "clicks",
         "sessions", "purchases", "revenue")  # fmt: skip
PURCHASES = ("user_pseudo_id", "session_id", "date", "channel_group", "paid", "revenue",
             "month", "first_purchase", "acquired_on", "cohort_month", "acquired_via",
             "months_since", "first_in_month", "days_since")  # fmt: skip


class WorkbookError(Exception):
    """The inputs cannot be read, or a workbook cannot be checked."""


def serial(day: dt.date) -> int:
    """A date as Excel's serial day number."""
    return (day - EPOCH).days


def from_serial(value) -> dt.date | None:
    """A cell's date, from a serial day number or the datetime openpyxl reads one as."""
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    return EPOCH + dt.timedelta(days=int(value))


@dataclass(frozen=True)
class Inputs:
    daily: list[dict]
    purchases: list[dict]


def read_inputs(conninfo: str) -> Inputs:
    """The rows the workbook takes as values, from the analysis views."""
    try:
        with psycopg.connect(conninfo, row_factory=dict_row) as conn:
            daily = conn.execute(
                "SELECT date, channel, channel_group, spend, impressions, clicks, sessions, "
                "purchases, revenue FROM analysis.channel_daily ORDER BY date, channel"
            ).fetchall()
            purchases = conn.execute(
                "SELECT user_pseudo_id, session_id, session_date AS date, channel_group, "
                "paid, revenue FROM analysis.purchases "
                'ORDER BY user_pseudo_id COLLATE "C", session_start, session_id'
            ).fetchall()
    except psycopg.Error as exc:
        raise WorkbookError(f"cannot read the analysis views: {exc}".strip()) from exc
    if not daily:
        raise WorkbookError("analysis.channel_daily has no rows: load the data first")
    return Inputs(daily, purchases)


def _months(first: dt.date, last: dt.date) -> list[dt.date]:
    out, m = [], first.replace(day=1)
    while m <= last:
        out.append(m)
        m = (m + dt.timedelta(days=32)).replace(day=1)
    return out


def _number(value):
    return float(value) if isinstance(value, Decimal) else value


class _Sheet:
    """A sheet written row by row, with headers in bold."""

    def __init__(self, wb: Workbook, title: str, headers: Iterable[str], start: int = 1):
        self.ws = wb.create_sheet(title)
        self.row = start
        self.append(list(headers))
        for cell in self.ws[start]:
            cell.font = Font(bold=True)
        self.ws.freeze_panes = self.ws.cell(start + 1, 1)

    def append(self, values: list, dates: Iterable[int] = ()) -> int:
        r = self.row
        for c, v in enumerate(values, 1):
            cell = self.ws.cell(r, c, _number(v))
            if c in dates:
                cell.number_format = DATE_FORMAT
        self.row += 1
        return r


def build(inputs: Inputs, path: Path, *, ltv_days: int = LTV_DAYS) -> dict[str, int]:
    """Write the workbook to ``path``; the number of rows on each sheet."""
    if isinstance(ltv_days, bool) or not isinstance(ltv_days, int) or ltv_days < 1:
        raise WorkbookError(f"ltv_days must be a whole number of days, at least 1: {ltv_days!r}")
    daily, purchases = inputs.daily, inputs.purchases
    first_day = min(r["date"] for r in daily)
    last_day = max(r["date"] for r in daily)
    months = _months(first_day, last_day)
    channels: dict[str, str] = {}
    for r in daily:
        channels.setdefault(r["channel"], r["channel_group"])
    paid_groups = sorted(set(channels.values()))
    first_groups = {p["channel_group"] for p in _firsts(purchases)}
    groups = paid_groups + sorted(first_groups - set(paid_groups))
    cohorts = sorted({p["date"].replace(day=1) for p in _firsts(purchases)})

    wb = Workbook()
    about = wb.active
    about.title = "About"
    nd, np_ = len(daily) + 1, len(purchases) + 1  # last data row on each input sheet

    def d(col: str) -> str:
        return f"Daily!${col}$2:${col}${nd}"

    def p(col: str) -> str:
        return f"Purchases!${col}$2:${col}${max(np_, 2)}"

    # Input columns: Daily A date, B month, C channel, D group, E spend, F impressions,
    # G clicks, H sessions, I purchases, J revenue.
    sheet = _Sheet(wb, "Daily", DAILY)
    for r in daily:
        n = sheet.row
        sheet.append(
            [serial(r["date"]), f"=DATE(YEAR(A{n}),MONTH(A{n}),1)", r["channel"],
             r["channel_group"], r["spend"], r["impressions"], r["clicks"], r["sessions"],
             r["purchases"], r["revenue"]],
            dates=(1, 2),
        )  # fmt: skip
    # Purchases A user, B session, C date, D group, E paid, F revenue; formulas G month,
    # H first purchase, I acquired on, J cohort month, K acquired via, L months since,
    # M first purchase in its month, N days since acquired.
    sheet = _Sheet(wb, "Purchases", PURCHASES)
    for r in purchases:
        n, m = sheet.row, sheet.row - 1
        sheet.append(
            [r["user_pseudo_id"], r["session_id"], serial(r["date"]), r["channel_group"],
             int(r["paid"]), r["revenue"],
             f"=DATE(YEAR(C{n}),MONTH(C{n}),1)",
             f"=IF(A{n}<>A{m},1,0)",
             f"=IF(H{n}=1,C{n},I{m})",
             f"=DATE(YEAR(I{n}),MONTH(I{n}),1)",
             f"=IF(H{n}=1,D{n},K{m})",
             f"=(YEAR(G{n})-YEAR(J{n}))*12+MONTH(G{n})-MONTH(J{n})",
             f"=IF(OR(A{n}<>A{m},G{n}<>G{m}),1,0)",
             f"=C{n}-I{n}"],
            dates=(3, 7, 9, 10),
        )  # fmt: skip

    about["A1"] = "Marketing ROI: the analysis views rebuilt in formulas"
    about["A1"].font = Font(bold=True, size=13)
    about["A3"], about["B3"] = "First day", f"=MIN({d('A')})"
    about["A4"], about["B4"] = "Last day", f"=MAX({d('A')})"
    for cell in ("B3", "B4"):
        about[cell].number_format = DATE_FORMAT
    notes = [
        "The data is synthetic: generated by marketing_roi.generate, not collected.",
        "Purchases are credited by last click, as GA4 records them; these are last-click "
        "figures with last click's biases.",
        "Daily and Purchases are inputs (values from the analysis views); every other "
        "figure is a formula over them. An empty cell is SQL's NULL.",
        "Channels = analysis.channel_performance; Acquisition = analysis.acquisition; "
        "Blended = analysis.blended; Cohorts = analysis.cohorts; LTV = "
        "analysis.ltv_by_channel(window), the window being LTV!B1.",
        "ROI is on revenue, (revenue - spend) / spend: there is no cost of goods.",
    ]
    for i, text in enumerate(notes, 6):
        about[f"A{i}"] = text
    about.column_dimensions["A"].width = 18
    last = "About!$B$4"

    # Channels: one row per channel and month, then the channel's whole period.
    sheet = _Sheet(wb, "Channels", ("month", "channel", "channel_group", "spend",
                   "impressions", "clicks", "sessions", "purchases", "revenue", "ctr", "cpc",
                   "conversion_rate", "cost_per_purchase", "roas"))  # fmt: skip
    for channel, group in sorted(channels.items()):
        for month in [*months, None]:
            n = sheet.row
            by = f",{d('B')},$A{n}" if month else ""
            sums = [f"=SUMIFS({d(col)},{d('C')},$B{n}{by})" for col in "EFGHIJ"]
            sheet.append(
                [serial(month) if month else None, channel, group, *sums,
                 f'=IF(E{n}=0,"",F{n}/E{n})', f'=IF(F{n}=0,"",D{n}/F{n})',
                 f'=IF(F{n}=0,"",H{n}/F{n})', f'=IF(H{n}=0,"",D{n}/H{n})',
                 f'=IF(D{n}=0,"",I{n}/D{n})'],
                dates=(1,),
            )  # fmt: skip

    # Acquisition: new customers by the group their first purchase is credited to.
    sheet = _Sheet(wb, "Acquisition", ("month", "channel_group", "paid", "spend",
                   "new_customers", "cac"))  # fmt: skip
    for group in groups:
        for month in [*months, None]:
            n = sheet.row
            spend_by = f",{d('B')},$A{n}" if month else ""
            new_by = f",{p('G')},$A{n}" if month else ""
            sheet.append(
                [serial(month) if month else None, group,
                 f"=COUNTIFS({d('D')},$B{n})>0",
                 f"=SUMIFS({d('E')},{d('D')},$B{n}{spend_by})",
                 f"=COUNTIFS({p('H')},1,{p('D')},$B{n}{new_by})",
                 f'=IF(C{n},IF(E{n}=0,"",D{n}/E{n}),"")'],
                dates=(1,),
            )  # fmt: skip

    # Blended: all spend against all revenue and new customers, by month and in all.
    sheet = _Sheet(wb, "Blended", ("month", "spend", "purchases", "revenue", "new_customers",
                   "new_customer_revenue", "paid_purchases", "paid_revenue", "blended_roas",
                   "blended_roi", "paid_roas", "blended_cac"))  # fmt: skip
    for month in [*months, None]:
        n = sheet.row
        if month:
            g = f"{p('G')},$A{n}"
            figures = [f"=SUMIFS({d('E')},{d('B')},$A{n})", f"=COUNTIFS({g})",
                       f"=SUMIFS({p('F')},{g})", f"=COUNTIFS({g},{p('H')},1)",
                       f"=SUMIFS({p('F')},{g},{p('H')},1)", f"=COUNTIFS({g},{p('E')},1)",
                       f"=SUMIFS({p('F')},{g},{p('E')},1)"]  # fmt: skip
        else:
            figures = [f"=SUM({d('E')})", f"=COUNT({p('F')})", f"=SUM({p('F')})",
                       f"=SUM({p('H')})", f"=SUMIFS({p('F')},{p('H')},1)", f"=SUM({p('E')})",
                       f"=SUMIFS({p('F')},{p('E')},1)"]  # fmt: skip
        sheet.append(
            [serial(month) if month else None, *figures,
             f'=IF(B{n}=0,"",D{n}/B{n})', f'=IF(B{n}=0,"",(D{n}-B{n})/B{n})',
             f'=IF(B{n}=0,"",H{n}/B{n})', f'=IF(E{n}=0,"",B{n}/E{n})'],
            dates=(1,),
        )  # fmt: skip

    # Cohorts: each first-purchase month, followed to the data's last month.
    sheet = _Sheet(wb, "Cohorts", ("cohort_month", "months_since", "month", "cohort_size",
                   "active_customers", "purchases", "revenue", "retention",
                   "cumulative_revenue_per_customer", "cumulative_purchases_per_customer",
                   "complete"))  # fmt: skip
    for cohort in cohorts:
        for k in range(len(_months(cohort, last_day))):
            n = sheet.row
            c = f"{p('J')},$A{n}"
            sheet.append(
                [serial(cohort), k, f"=DATE(YEAR(A{n}),MONTH(A{n})+B{n},1)",
                 f"=COUNTIFS({p('H')},1,{c})",
                 f"=COUNTIFS({c},{p('G')},$C{n},{p('M')},1)",
                 f"=COUNTIFS({c},{p('G')},$C{n})",
                 f"=SUMIFS({p('F')},{c},{p('G')},$C{n})",
                 f"=E{n}/D{n}",
                 f'=SUMIFS({p("F")},{c},{p("L")},"<="&$B{n})/D{n}',
                 f'=COUNTIFS({c},{p("L")},"<="&$B{n})/D{n}',
                 f"=DATE(YEAR(C{n}),MONTH(C{n})+1,1)-1<={last}"],
                dates=(1, 3),
            )  # fmt: skip

    # LTV: revenue per customer over the window from their first purchase, against CAC
    # over the same acquisition days; only customers followed for the whole window.
    sheet = _Sheet(wb, "LTV", ("channel_group", "paid", "customers", "spend", "cac",
                   "revenue_per_customer", "ltv_to_cac"), start=4)  # fmt: skip
    ws = sheet.ws
    ws["A1"], ws["B1"] = "Window (days)", ltv_days
    ws["A2"], ws["B2"] = "Last acquisition day", f"={last}-$B$1+1"
    ws["B2"].number_format = DATE_FORMAT
    ws["C1"] = "an input: change it and the sheet follows"
    for group in sorted(groups, key=lambda g: (g not in paid_groups, g)):
        n = sheet.row
        sheet.append(
            [group, f"=COUNTIFS({d('D')},$A{n})>0",
             f'=COUNTIFS({p("H")},1,{p("D")},$A{n},{p("C")},"<="&$B$2)',
             f'=IF(B{n},SUMIFS({d("E")},{d("D")},$A{n},{d("A")},"<="&$B$2),"")',
             f'=IF(OR(C{n}=0,NOT(B{n})),"",D{n}/C{n})',
             f'=IF(C{n}=0,"",SUMIFS({p("F")},{p("K")},$A{n},{p("I")},"<="&$B$2,'
             f'{p("N")},"<"&$B$1)/C{n})',
             f'=IF(OR(C{n}=0,NOT(B{n})),"",IF(D{n}=0,"",F{n}/E{n}))'],
        )  # fmt: skip

    for ws in wb.worksheets:
        if ws.title != "About":
            for i in range(1, ws.max_column + 1):
                ws.column_dimensions[get_column_letter(i)].width = 14
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {ws.title: ws.max_row for ws in wb.worksheets}


def _firsts(purchases: list[dict]) -> list[dict]:
    """Each customer's first purchase, from rows sorted by customer and time."""
    out, previous = [], object()
    for r in purchases:
        if r["user_pseudo_id"] != previous:
            out.append(r)
        previous = r["user_pseudo_id"]
    return out


# --- checking a calculated workbook against the views ---------------------------------------

Value = Callable[[str, int, int], object]
"""A calculated cell's value: (sheet, row, column) -> value."""

SHEETS = {
    "Channels": (("month", "channel"), 4),
    "Acquisition": (("month", "channel_group"), 3),
    "Blended": (("month",), 2),
    "Cohorts": (("cohort_month", "months_since"), 3),
}
"""Per output sheet: the key columns, and the first column of figures (1-based)."""
VIEWS = {
    "Channels": "channel_performance",
    "Acquisition": "acquisition",
    "Blended": "blended",
    "Cohorts": "cohorts",
}
DATE_COLUMNS = {"month", "cohort_month"}


def _same(cell, sql) -> bool:
    if sql is None or cell is None or cell == "":
        return (sql is None) == (cell is None or cell == "")
    if isinstance(sql, bool) or isinstance(cell, bool):
        return bool(cell) == bool(sql) and isinstance(cell, bool | int | float)
    if isinstance(sql, dt.date):
        return from_serial(cell) == sql
    if isinstance(sql, int | float | Decimal):
        if isinstance(cell, str):
            return False
        a, b = float(cell), float(sql)
        return math.isclose(a, b, rel_tol=1e-9, abs_tol=1e-9)
    return str(cell) == str(sql)


def check(conninfo: str, headers: dict[str, list[str]], value: Value) -> dict:
    """Compare every figure on the output sheets with the views it rebuilds.

    ``headers`` holds each sheet's header row and ``value`` gives a calculated cell, so
    the same check runs on a workbook a spreadsheet program has calculated and on one a
    test evaluated. Returns the number of figures compared and the mismatches.
    """
    compared, wrong = 0, []
    try:
        with psycopg.connect(conninfo, row_factory=dict_row) as conn:
            sql = {s: conn.execute(f"SELECT * FROM analysis.{v}").fetchall()
                   for s, v in VIEWS.items()}  # fmt: skip
            window = value("LTV", 1, 2)
            if isinstance(window, bool) or not isinstance(window, int | float) or window < 1:
                raise WorkbookError(f"LTV!B1 must be a number of days, got {window!r}")
            sql["LTV"] = conn.execute(
                "SELECT * FROM analysis.ltv_by_channel(%s)", (int(window),)
            ).fetchall()
    except psycopg.Error as exc:
        raise WorkbookError(f"cannot read the analysis views: {exc}".strip()) from exc
    layouts = dict(SHEETS, LTV=(("channel_group",), 2))
    for name, (keys, first) in layouts.items():
        cols = headers[name]
        top = 5 if name == "LTV" else 2
        by_key = {tuple(r[k] for k in keys): r for r in sql[name]}
        seen = set()
        r = top
        while value(name, r, 1) is not None or (name != "LTV" and value(name, r, 2) is not None):
            key = tuple(
                from_serial(value(name, r, cols.index(k) + 1))
                if k in DATE_COLUMNS
                else value(name, r, cols.index(k) + 1)
                for k in keys
            )
            row = by_key.get(key)
            if row is None:
                # A row the view leaves out has nothing in it: no spend and no customers
                # (Acquisition), or no customers followed for the whole window (LTV).
                empty = ("spend", "new_customers") if name == "Acquisition" else ("customers",)
                for k in empty:
                    compared += 1
                    v = value(name, r, cols.index(k) + 1)
                    if v not in (0, 0.0, "", None):
                        wrong.append((name, r, k, v, "row not in the view"))
            else:
                seen.add(key)
                for c in range(first, len(cols) + 1):
                    k = cols[c - 1]
                    compared += 1
                    v = value(name, r, c)
                    if not _same(v, row[k]):
                        wrong.append((name, r, k, v, row[k]))
            r += 1
        for key in by_key.keys() - seen:
            wrong.append((name, None, key, None, "view row missing from the sheet"))
    return {"compared": compared, "wrong": wrong}


def check_file(conninfo: str, path: Path) -> dict:
    """``check`` on a workbook saved with its calculated values."""
    wb = load_workbook(path, data_only=True, read_only=False)
    names = [*SHEETS, "LTV"]
    missing = [n for n in names if n not in wb.sheetnames]
    if missing:
        raise WorkbookError(f"{path} has no {', '.join(missing)} sheet")
    headers = {n: [c.value for c in wb[n][4 if n == "LTV" else 1]] for n in names}
    if wb["Channels"].cell(2, 4).value is None:
        raise WorkbookError(
            f"{path} holds formulas without calculated values: open it in a spreadsheet "
            "program and save it, or run scripts/excel_recalc.ps1 on it, first"
        )

    def value(sheet: str, row: int, col: int):
        return wb[sheet].cell(row, col).value

    return check(conninfo, headers, value)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="python -m marketing_roi.workbook",
        description="Write the Excel workbook that rebuilds the analysis views in formulas, "
        "or check a calculated one against them.",
    )
    parser.add_argument("--dsn", default=os.environ.get("MARKETING_ROI_DSN"),
                        help="connection string (default: $MARKETING_ROI_DSN)")  # fmt: skip
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--out", type=Path, help="write the workbook here")
    group.add_argument("--check", type=Path, help="compare a calculated workbook with the views")
    parser.add_argument("--ltv-days", type=int, default=LTV_DAYS, help="the LTV sheet's window")
    args = parser.parse_args(argv)
    if not args.dsn:
        parser.error("no database: pass --dsn or set MARKETING_ROI_DSN")
    if args.ltv_days < 1:
        parser.error("--ltv-days must be at least 1")
    try:
        if args.out:
            rows = build(read_inputs(args.dsn), args.out, ltv_days=args.ltv_days)
            print(f"wrote {args.out}: " + ", ".join(f"{k} {v:,} rows" for k, v in rows.items()))
            return 0
        result = check_file(args.dsn, args.check)
    except WorkbookError as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1
    for sheet, row, column, got, want in result["wrong"][:20]:
        print(f"  {sheet} row {row} {column}: workbook {got!r}, view {want!r}")
    print(f"{result['compared']:,} figures compared, {len(result['wrong'])} differ")
    return 1 if result["wrong"] else 0


if __name__ == "__main__":
    sys.exit(main())
